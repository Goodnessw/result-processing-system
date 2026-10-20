from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core import signing
from django.core.exceptions import ValidationError
from django.db import transaction
from django.http import HttpResponse, HttpResponseForbidden
from django.shortcuts import get_object_or_404, redirect, render
from django.views.decorators.http import require_http_methods

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from accounts.models import User
from academics.models import Course, CourseRegistration
from .models import Result, ResultAuditLog


TEMPLATE_SALT = "results.excel_template"
PREVIEW_SESSION_KEY = "result_excel_preview"
MAX_UPLOAD_SIZE = 5 * 1024 * 1024

HEADERS = [
    "Matric Number",
    "Student Name",
    "CA Score (0-30)",
    "Exam Score (0-70)",
]


def get_lecturer_course(request, course_id):
    if request.user.role != User.Role.LECTURER:
        return None

    return Course.objects.filter(
        pk=course_id,
        lecturer=request.user,
        is_active=True,
    ).select_related(
        "programme",
        "semester",
        "semester__session",
    ).first()


def get_registered_students(course):
    return list(
        CourseRegistration.objects.filter(
            course=course,
            status=CourseRegistration.RegistrationStatus.REGISTERED,
        )
        .select_related("student__user")
        .order_by(
            "student__user__last_name",
            "student__user__first_name",
        )
    )


def get_matric_number(registration):
    return (registration.student.user.matric_number or "").strip()


def get_student_name(registration):
    user = registration.student.user
    return user.get_full_name().strip() or user.username


@login_required
def download_result_template(request, course_id):
    course = get_lecturer_course(request, course_id)

    if course is None:
        if request.user.role != User.Role.LECTURER:
            return HttpResponseForbidden("Lecturer access required.")
        return HttpResponseForbidden(
            "This course is not assigned to you or is inactive."
        )

    registrations = get_registered_students(course)

    if not registrations:
        messages.error(
            request,
            "There are no registered students to include in the template.",
        )
        return redirect("course_results", course_id=course.id)

    matric_numbers = [get_matric_number(r) for r in registrations]

    if any(not matric for matric in matric_numbers):
        messages.error(
            request,
            "Every registered student must have a matric number before "
            "an Excel template can be generated.",
        )
        return redirect("course_results", course_id=course.id)

    if len(set(matric_numbers)) != len(matric_numbers):
        messages.error(
            request,
            "Duplicate matric numbers were found. Correct the student records first.",
        )
        return redirect("course_results", course_id=course.id)

    payload = {
        "version": 1,
        "course_id": course.id,
        "registration_ids": [r.id for r in registrations],
        "matric_numbers": matric_numbers,
    }

    template_token = signing.dumps(payload, salt=TEMPLATE_SALT)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Result Entry"

    sheet.merge_cells("A1:D1")
    sheet["A1"] = f"{course.code} — RESULT ENTRY TEMPLATE"
    sheet["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    sheet["A1"].fill = PatternFill(
        "solid", fgColor="244FC5"
    )
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 30

    sheet.merge_cells("A2:D2")
    sheet["A2"] = f"Course: {course.title} ({course.code})"

    sheet.merge_cells("A3:D3")
    sheet["A3"] = (
        f"Session: {course.semester.session} | "
        f"Semester: {course.semester.get_name_display()}"
    )

    sheet.merge_cells("A4:D4")
    sheet["A4"] = (
        "Enter CA scores from 0 to 30 and Exam scores from 0 to 70. "
        "Do not change student details or column headings."
    )
    sheet["A4"].alignment = Alignment(wrap_text=True)
    sheet.row_dimensions[4].height = 32

    for column, heading in enumerate(HEADERS, start=1):
        cell = sheet.cell(row=5, column=column, value=heading)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="334155")
        cell.alignment = Alignment(horizontal="center")

    for row_number, registration in enumerate(registrations, start=6):
        sheet.cell(row_number, 1, get_matric_number(registration))
        sheet.cell(row_number, 2, get_student_name(registration))
        sheet.cell(row_number, 3, None)
        sheet.cell(row_number, 4, None)

        sheet.cell(row_number, 1).number_format = "@"
        sheet.cell(row_number, 3).number_format = "0.00"
        sheet.cell(row_number, 4).number_format = "0.00"

    widths = [24, 32, 20, 20]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    sheet.freeze_panes = "A6"
    sheet.auto_filter.ref = f"A5:D{5 + len(registrations)}"

    system_sheet = workbook.create_sheet("_SYSTEM")
    system_sheet["A1"] = template_token
    system_sheet.sheet_state = "hidden"

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )
    response["Content-Disposition"] = (
        f'attachment; filename="{course.code}_result_template.xlsx"'
    )
    return response


@login_required
@require_http_methods(["POST"])
def upload_result_excel(request, course_id):
    course = get_lecturer_course(request, course_id)

    if course is None:
        if request.user.role != User.Role.LECTURER:
            return HttpResponseForbidden("Lecturer access required.")
        return HttpResponseForbidden(
            "This course is not assigned to you or is inactive."
        )

    uploaded_file = request.FILES.get("excel_file")

    if not uploaded_file:
        messages.error(request, "Choose the completed Excel template first.")
        return redirect("course_results", course_id=course.id)

    if uploaded_file.size > MAX_UPLOAD_SIZE:
        messages.error(request, "The Excel file must not exceed 5 MB.")
        return redirect("course_results", course_id=course.id)

    if not uploaded_file.name.lower().endswith(".xlsx"):
        messages.error(
            request,
            "Only the system-generated .xlsx template is supported.",
        )
        return redirect("course_results", course_id=course.id)

    try:
        workbook = load_workbook(
            uploaded_file,
            read_only=False,
            data_only=False,
        )

        if (
            "Result Entry" not in workbook.sheetnames
            or "_SYSTEM" not in workbook.sheetnames
        ):
            raise ValueError(
                "This is not a recognized system-generated template."
            )

        token = workbook["_SYSTEM"]["A1"].value

        if not isinstance(token, str) or not token:
            raise ValueError("The template identifier is missing.")

        payload = signing.loads(token, salt=TEMPLATE_SALT)

        if (
            payload.get("version") != 1
            or payload.get("course_id") != course.id
        ):
            raise ValueError(
                "This template was not generated for the selected course."
            )

        registrations = get_registered_students(course)
        registration_ids = [r.id for r in registrations]
        matric_numbers = [get_matric_number(r) for r in registrations]

        if (
            payload.get("registration_ids") != registration_ids
            or payload.get("matric_numbers") != matric_numbers
        ):
            raise ValueError(
                "The registered student list has changed since the template "
                "was downloaded. Download a fresh template."
            )

        if len(set(matric_numbers)) != len(matric_numbers):
            raise ValueError(
                "Duplicate matric numbers exist in the registered student list."
            )

        sheet = workbook["Result Entry"]

        actual_headers = [
            sheet.cell(5, column).value
            for column in range(1, 5)
        ]

        if actual_headers != HEADERS:
            raise ValueError(
                "The template headings or column order have been changed."
            )

        registration_by_matric = {
            get_matric_number(r): r for r in registrations
        }

        uploaded_rows = {}
        errors = []

        for row_number in range(6, sheet.max_row + 1):
            values = [
                sheet.cell(row_number, column).value
                for column in range(1, 5)
            ]

            if all(value is None for value in values):
                continue

            if any(
                isinstance(value, str) and value.startswith("=")
                for value in values
            ):
                errors.append(
                    f"Row {row_number}: formulas are not allowed."
                )
                continue

            matric_value, name_value, ca_value, exam_value = values

            if not isinstance(matric_value, str):
                errors.append(
                    f"Row {row_number}: matric number must remain text."
                )
                continue

            matric = matric_value.strip()

            if matric not in registration_by_matric:
                errors.append(
                    f"Row {row_number}: unknown or unregistered matric number "
                    f"'{matric}'."
                )
                continue

            if matric in uploaded_rows:
                errors.append(
                    f"Row {row_number}: duplicate matric number '{matric}'."
                )
                continue

            registration = registration_by_matric[matric]

            if (
                not isinstance(name_value, str)
                or name_value.strip() != get_student_name(registration)
            ):
                errors.append(
                    f"Row {row_number}: the name for {matric} was changed."
                )
                continue

            try:
                if ca_value is None or exam_value is None:
                    raise InvalidOperation

                ca_score = Decimal(str(ca_value))
                exam_score = Decimal(str(exam_value))

                if not ca_score.is_finite() or not exam_score.is_finite():
                    raise InvalidOperation

            except (InvalidOperation, ValueError, TypeError):
                errors.append(
                    f"Row {row_number} ({matric}): enter both CA and Exam scores."
                )
                continue

            if (
                ca_score < Decimal("0")
                or ca_score > Decimal("30")
                or exam_score < Decimal("0")
                or exam_score > Decimal("70")
            ):
                errors.append(
                    f"Row {row_number} ({matric}): CA must be 0–30 "
                    "and Exam must be 0–70."
                )
                continue

            if (
                ca_score.as_tuple().exponent < -2
                or exam_score.as_tuple().exponent < -2
            ):
                errors.append(
                    f"Row {row_number} ({matric}): scores may have at most "
                    "two decimal places."
                )
                continue

            uploaded_rows[matric] = {
                "registration_id": registration.id,
                "matric_number": matric,
                "student_name": get_student_name(registration),
                "ca_score": str(ca_score),
                "exam_score": str(exam_score),
                "total_score": str(ca_score + exam_score),
            }

        missing = set(registration_by_matric) - set(uploaded_rows)

        for matric in sorted(missing):
            errors.append(f"Missing result for {matric}.")

        if errors:
            for error in errors[:20]:
                messages.error(request, error)

            if len(errors) > 20:
                messages.error(
                    request,
                    f"{len(errors) - 20} additional validation errors were found.",
                )

            return redirect("course_results", course_id=course.id)

        existing_results = Result.objects.filter(
            course_registration__in=registrations
        )

        if existing_results.filter(
            status__in=[
                Result.Status.SUBMITTED,
                Result.Status.APPROVED,
            ]
        ).exists():
            messages.error(
                request,
                "This course contains submitted or approved results. "
                "The spreadsheet cannot overwrite locked results.",
            )
            return redirect("course_results", course_id=course.id)

        request.session[PREVIEW_SESSION_KEY] = {
            "course_id": course.id,
            "rows": list(uploaded_rows.values()),
        }

        return redirect("preview_result_excel", course_id=course.id)

    except (
        ValueError,
        KeyError,
        TypeError,
        signing.BadSignature,
        signing.SignatureExpired,
        OSError,
    ) as exc:
        messages.error(
            request,
            f"Excel upload rejected: {exc}",
        )
        return redirect("course_results", course_id=course.id)

    finally:
        if "workbook" in locals():
            workbook.close()


@login_required
@require_http_methods(["GET", "POST"])
def preview_result_excel(request, course_id):
    course = get_lecturer_course(request, course_id)

    if course is None:
        if request.user.role != User.Role.LECTURER:
            return HttpResponseForbidden("Lecturer access required.")
        return HttpResponseForbidden(
            "This course is not assigned to you or is inactive."
        )

    preview = request.session.get(PREVIEW_SESSION_KEY)

    if not preview or preview.get("course_id") != course.id:
        messages.error(
            request,
            "Your upload preview has expired. Upload the spreadsheet again.",
        )
        return redirect("course_results", course_id=course.id)

    if request.method == "GET":
        return render(
            request,
            "results/excel_preview.html",
            {
                "course": course,
                "rows": preview["rows"],
                "total_students": len(preview["rows"]),
                "dashboard_url": "/",
            },
        )

    action = request.POST.get("action")

    if action not in ("save", "submit"):
        messages.error(request, "Invalid action.")
        return redirect("preview_result_excel", course_id=course.id)

    rows = preview["rows"]
    registrations = get_registered_students(course)
    registrations_by_id = {r.id: r for r in registrations}

    if (
        len(rows) != len(registrations)
        or {row["registration_id"] for row in rows}
        != set(registrations_by_id)
    ):
        request.session.pop(PREVIEW_SESSION_KEY, None)
        messages.error(
            request,
            "The registered student list has changed. Download a new template.",
        )
        return redirect("course_results", course_id=course.id)

    if any(
        result and result.status in (
            Result.Status.SUBMITTED,
            Result.Status.APPROVED,
        )
        for result in Result.objects.filter(
            course_registration__course=course,
        )
    ):
        messages.info(
            request,
            "This course has already been submitted or approved. No new uploads can be processed.",
        )
        return redirect("course_results", course_id=course.id)

    try:
        with transaction.atomic():
            existing_results = {
                result.course_registration_id: result
                for result in Result.objects.select_for_update().filter(
                    course_registration__in=registrations
                )
            }

            for result in existing_results.values():
                if result.status in (
                    Result.Status.SUBMITTED,
                    Result.Status.APPROVED,
                ):
                    raise ValueError(
                        "A result was submitted or approved while this "
                        "spreadsheet was being reviewed. Nothing was saved."
                    )

            for row in rows:
                registration = registrations_by_id[row["registration_id"]]
                result = existing_results.get(registration.id)
                is_new = result is None

                if is_new:
                    result = Result(
                        course_registration=registration,
                        entered_by=request.user,
                    )

                result.ca_score = Decimal(row["ca_score"])
                result.exam_score = Decimal(row["exam_score"])
                result.entered_by = request.user
                result.approved_by = None
                result.status = (
                    Result.Status.SUBMITTED
                    if action == "submit"
                    else Result.Status.DRAFT
                )

                result.full_clean()
                result.save()

                if action == "save":
                    ResultAuditLog.objects.create(
                        result=result,
                        action=(
                            ResultAuditLog.Action.CREATED
                            if is_new
                            else ResultAuditLog.Action.UPDATED
                        ),
                        performed_by=request.user,
                        details="Result imported from the system Excel template.",
                    )
                else:
                    if is_new:
                        ResultAuditLog.objects.create(
                            result=result,
                            action=ResultAuditLog.Action.CREATED,
                            performed_by=request.user,
                            details="Result imported from the system Excel template.",
                        )

                    ResultAuditLog.objects.create(
                        result=result,
                        action=ResultAuditLog.Action.SUBMITTED,
                        performed_by=request.user,
                        details="Excel-imported result submitted for approval.",
                    )

    except (ValidationError, ValueError) as exc:
        messages.error(request, f"Import failed: {exc}")
        return redirect("preview_result_excel", course_id=course.id)

    request.session.pop(PREVIEW_SESSION_KEY, None)

    if action == "submit":
        messages.success(
            request,
            "All results were imported and submitted for approval.",
        )
    else:
        messages.success(
            request,
            "All results were imported and saved as drafts.",
        )

    return redirect("course_results", course_id=course.id)